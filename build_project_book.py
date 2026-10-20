# -*- coding: utf-8 -*-
r"""
生成《智能考场与AI结合的先进化终端》参赛作品项目书（修订版）。

内容依据 2026-10-20 当前项目树：
  00 集成联动、01 人脸、02 摄像头、03 终端、05 数据保护、06 安检、
  07 无线信号、08 试卷/保密室/火检、09 无障碍、10 考场模板、11 局域网上传。
  04 巡考机器人仍为方案文档；原 11/13 目录已删除/合并。
"""
from __future__ import annotations

import os
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "智能考场与AI结合的先进化终端_参赛作品项目书_修订版.docx")
ARCH = os.path.join(HERE, "_project_arch.png")

NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK = RGBColor(0x20, 0x20, 0x20)
GRAY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = "DEEAF6"
LIGHTER = "F2F7FB"


# ---------------------------------------------------------------------------
# 基础排版
# ---------------------------------------------------------------------------
def set_run_font(run, name="宋体", size=12, bold=None, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def add_rich(paragraph, text, size=12, color=DARK, base_bold=False):
    """支持 **加粗** 的极简行内格式。"""
    parts = str(text).split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_run_font(run, "宋体", size, base_bold or (i % 2 == 1), color)


def body(text, indent=True, space_after=4):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.first_line_indent = Pt(24) if indent else Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    add_rich(p, text)
    return p


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.75 + 0.55 * level)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(2)
    add_rich(p, text)
    return p


def h1(text):
    p = doc.add_paragraph(style="Heading 1")
    add_rich(p, text, 16, NAVY, True)
    p.paragraph_format.first_line_indent = Pt(0)
    return p


def h2(text):
    p = doc.add_paragraph(style="Heading 2")
    add_rich(p, text, 13.5, BLUE, True)
    p.paragraph_format.first_line_indent = Pt(0)
    return p


def h3(text):
    p = doc.add_paragraph(style="Heading 3")
    add_rich(p, text, 12, DARK, True)
    p.paragraph_format.first_line_indent = Pt(0)
    return p


def caption(text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, text, 9, GRAY)
    return p


def page_break():
    doc.add_page_break()


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def cell_margins(table, top=70, bottom=70, left=100, right=100):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + name)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def table(headers, rows, widths=None, font_size=9, align_center=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t)
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        shade(hdr[i], "1F4E79")
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_rich(p, htext, font_size, WHITE, True)
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    repeat_header(t.rows[0])
    no_split(t.rows[0])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        no_split(t.rows[-1])
        for ci, text in enumerate(row):
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            if align_center and ci in align_center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich(p, str(text), font_size)
            cells[ci].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ri % 2 == 1:
                shade(cells[ci], LIGHTER)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def add_field(paragraph, instr, placeholder=""):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_sep)
    run._r.append(txt)
    run._r.append(fld_end)
    set_run_font(run, "宋体", 9, False, GRAY)


def add_toc():
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "目录将在打开文档时自动更新（如未更新：右键 → 更新域）"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(txt)
    run._r.append(fld_end)
    set_run_font(run, "宋体", 10, False, GRAY)


# ---------------------------------------------------------------------------
# 架构示意图
# ---------------------------------------------------------------------------
def make_arch_image(path):
    from PIL import Image, ImageDraw, ImageFont

    W, H = 2000, 1120
    img = Image.new("RGB", (W, H), "#F7FAFD")
    d = ImageDraw.Draw(img)

    def font(size, bold=False):
        f = "C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"
        return ImageFont.truetype(f, size)

    def rr(box, fill, outline="#1F4E79", width=3, radius=24):
        d.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)

    def center(text, y, size, fill="#1F4E79", bold=True):
        f = font(size, bold)
        w = d.textbbox((0, 0), text, font=f)[2]
        d.text(((W - w) / 2, y), text, font=f, fill=fill)

    center("智能考场 AI 巡考系统 · 端-边-云协同架构", 28, 46)
    center("统一入口 · 本地推理 · 分级告警 · 人工复核 · 全程留痕", 92, 26, "#4A6B8A", False)

    boxes = [
        (60, 170, 560, 780, "端侧 · 考场现场", "#EAF3FB", [
            "摄像机 / 拾音器",
            "人脸核验终端",
            "安检门 / 手机探测",
            "无线频谱 / 屏蔽设备",
            "门禁 / 保密室传感器",
            "本机摄像头 / 视频文件",
        ]),
        (720, 170, 1280, 780, "边缘侧 · 考点本地", "#E8F6F1", [
            "02 画质诊断 / 行为识别",
            "01 人脸核验 / 证件核对",
            "06 安检判据 / 保管柜",
            "07 频谱监测 / 屏蔽管控",
            "08 试卷流转 / 保密室 / 火检",
            "10 普通 + 特殊考场模板",
            "本地 SQLite / 审计链 / 断网续传",
        ]),
        (1440, 170, 1940, 780, "业务侧 · 总控台 / 终端", "#F4EEF9", [
            "总控台：一次登录 + SSO",
            "03 告警处置 / 复核 / 驾驶舱",
            "05 数据保护 / 授权 / 审计",
            "11 局域网上传 / 终端回显",
            "00 集成联动总线",
            "单文件 exe / 源码树独立运行",
        ]),
    ]
    for x1, y1, x2, y2, title, fill, items in boxes:
        rr((x1, y1, x2, y2), fill)
        d.rectangle((x1, y1, x2, y1 + 78), fill="#1F4E79")
        f = font(30, True)
        tw = d.textbbox((0, 0), title, font=f)[2]
        d.text((x1 + (x2 - x1 - tw) / 2, y1 + 18), title, font=f, fill="white")
        f2 = font(25, False)
        yy = y1 + 108
        for item in items:
            d.ellipse((x1 + 34, yy + 12, x1 + 48, yy + 26), fill="#2E74B5")
            d.text((x1 + 62, yy), item, font=f2, fill="#203040")
            yy += 72

    def arrow(x1, y1, x2, y2, color="#2E74B5"):
        d.line((x1, y1, x2, y2), fill=color, width=9)
        d.polygon([(x2, y2), (x2 - 22, y2 - 13), (x2 - 22, y2 + 13)], fill=color)

    arrow(565, 475, 710, 475)
    arrow(1285, 475, 1430, 475)
    f = font(22, False)
    d.text((575, 425), "音视频/传感", font=f, fill="#4A6B8A")
    d.text((1292, 425), "告警/结构化数据", font=f, fill="#4A6B8A")

    rr((60, 850, 1940, 1045), "#FFFFFF", "#B9D3EA", 2, 18)
    f = font(25, True)
    d.text((90, 875), "交付与运行", font=f, fill="#1F4E79")
    f2 = font(23, False)
    lines = [
        "① 单文件 exe：启动器 + payload + 自带 Python 3.12，双击即用，约 511 MB",
        "② 源码树：项目树/All 下每个模块可独立运行；总控台无 runtime 时自动用系统 Python",
        "③ 一键打包：Git Bash 执行 一键打包.sh；图标取 icon/icon.png，资源嵌入启动器",
    ]
    yy = 918
    for line in lines:
        d.text((95, yy), line, font=f2, fill="#203040")
        yy += 40
    img.save(path)


# ---------------------------------------------------------------------------
# 文档初始化
# ---------------------------------------------------------------------------
def init_doc():
    global doc
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.4)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.font.color.rgb = DARK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)

    body_style = styles["Body Text"]
    body_style.font.name = "宋体"
    body_style.font.size = Pt(12)
    body_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    body_style.paragraph_format.line_spacing = 1.5
    body_style.paragraph_format.space_after = Pt(4)

    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13.5, BLUE),
                              ("Heading 3", 12, DARK)):
        st = styles[name]
        st.font.name = "微软雅黑"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.2

    cap = styles["Caption"]
    cap.font.name = "宋体"
    cap.font.size = Pt(9)
    cap.font.color.rgb = GRAY

    lb = styles["List Bullet"]
    lb.font.name = "宋体"
    lb.font.size = Pt(12)
    lb._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_rich(hp, "智能考场 AI 巡考系统 · 参赛作品项目书", 8, GRAY)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(fp, "第 ", 9, GRAY)
    add_field(fp, " PAGE ")
    add_rich(fp, " 页 / 共 ", 9, GRAY)
    add_field(fp, " NUMPAGES ")
    add_rich(fp, " 页", 9, GRAY)

    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


# ---------------------------------------------------------------------------
# 封面 / 声明 / 摘要 / 目录
# ---------------------------------------------------------------------------
def cover():
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, "智能考场 AI 巡考系统", 28, NAVY, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, "面向教育考试与日常教学场景的端-边-云智能巡考终端", 15, BLUE, True)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, "参赛作品项目书", 20, DARK, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, "（修订版 V2.0）", 12, GRAY)

    doc.add_paragraph()
    doc.add_paragraph()
    info = [
        ("作品类别", "科技发明制作 / 智能系统与软件平台"),
        ("参赛赛道 / 组别", "【待填】"),
        ("参赛团队名称", "【待填】"),
        ("团队成员", "【待填：姓名（专业 / 年级）】"),
        ("指导教师", "【待填：姓名・职称】"),
        ("所属单位", "【待填：院校 / 学院】"),
        ("提交日期", "2026 年【×】月"),
        ("版本 / 日期", "V2.0 / 2026-10-20"),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for k, v in info:
        cells = t.add_row().cells
        cells[0].width = Cm(4.2)
        cells[1].width = Cm(9.2)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_rich(p0, k + "：", 12, GRAY, True)
        p1 = cells[1].paragraphs[0]
        add_rich(p1, v, 12, DARK)
        for c in cells:
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    page_break()


def declaration():
    h1("原创性声明")
    body("本作品为团队独立原创设计，无抄袭、剽窃行为。系统架构、软件实现、测试方法与场景方案均为自主完成；"
         "公开资料、开源组件与第三方模型均在文中或代码仓库中标注来源。")
    body("方案遵循《数据安全法》《个人信息保护法》《网络安全法》及教育考试相关管理规定，"
         "文档不包含任何真实师生隐私数据；演示数据均标注“演示”，不冒充真实接入。")
    body("项目书所述“已实现”内容均可在项目树中运行复现，并附自检/测试证据；"
         "尚未完成的内容（如巡考机器人硬件、真实频谱仪/门禁/摄像头现场联调）已明确标注边界，"
         "不以预期代替成果。")
    doc.add_paragraph()
    body("团队成员签字：____________________　　指导教师签字：____________________", indent=False)
    body("日期：2026 年 ____ 月 ____ 日", indent=False)
    page_break()


def abstract():
    h1("摘要")
    body("本项目面向教育考试与日常教学场景，构建一套以“本地推理、分级告警、人工复核、全程留痕”为核心的"
         "端-边-云智能巡考系统。系统利旧考场既有摄像机、拾音器与门禁设备，在考点本地完成人脸核验、"
         "画质诊断、行为识别、智能安检、无线信号监测、试卷保密室监控与火场识别；仅将告警与结构化数据"
         "汇聚到智能终端，由人工完成确认、复核与处置，形成可追溯的闭环。")
    body("当前版本已形成可运行、可测试、可独立部署的软件原型：总控台统一登录一次后通过 SSO 令牌直达"
         "各模块；02 摄像头核查支持机位新增/编辑、基准图、三级巡检与整改工单；03 智能终端提供告警处置、"
         "复核留痕、驾驶舱与命令台；05 考生数据保护提供分级、加密、脱敏、销毁与审计；06 智能安检、"
         "07 无线信号、08 试卷/保密室/火检、09 无障碍、10 普通/特殊考场模板、11 局域网上传均已有实现。"
         "系统同时提供单文件 exe 与源码树两种交付形态，源码树下每个模块可独立运行。")
    p = doc.add_paragraph()
    add_rich(p, "关键词：", 12, DARK, True)
    add_rich(p, "智能巡考；端-边-云；分级告警；单点登录；考生数据保护；局域网接入；火场识别", 12, DARK)
    page_break()


def revision_note():
    h1("修订说明（相对原稿）")
    body("原稿以方案设想与“预期指标”为主，缺少可运行证据，且目录中存在重复专题。"
         "本次根据项目实际完成情况重写全文，重点变化如下：")
    table(
        ["原稿问题", "本次修订"],
        [
            ("以“预期/设计目标”为主，未区分已实现与规划", "逐项标注“已实现 / 部分实现 / 方案文档”，并以自检、测试与代码入口作为证据"),
            ("11 摄像头 AI 识别、13 火场识别为重复入口", "删除 11 与 13 两个目录；火检能力并入 08，摄像头 AI 能力由 02 覆盖，结构目录重排为 00–11"),
            ("未体现统一入口与模块间免重复登录", "总控台统一登录一次，01/02/03/05/08 通过 HMAC SSO 令牌直达，模块界面统一深色主题"),
            ("02 机位配置易失败、保存按钮被裁掉", "新增/编辑机位改为底部固定按钮，补安装位置、必填校验、异常提示与双击编辑"),
            ("03 告警列表整行红/黄底，文字对比度低", "改为深色行 + 级别文字上色/加粗，增加统计条、详情区与列宽调整"),
            ("缺少跨设备告警上传入口", "新增 11 告警上传 UI 与 0.0.0.0:5030 局域网 HTTP 接收，手机/其他电脑可上传到 03 并回显"),
            ("测试结果多为“预计通过”", "改为可复现的真实结果：02 108/108、08 114/114、00 39/39、05 174+110+16+6 等"),
        ],
        widths=[6.6, 10.6], font_size=9,
    )
    caption("表 0  本次修订与原稿的主要差异")
    page_break()


def toc():
    h1("目录")
    table(
        ["章节", "主要内容"],
        [
            ("摘要", "项目定位、已实现能力、交付形态与关键词"),
            ("修订说明", "相对原稿的变化、删除重复目录与真实测试证据"),
            ("一、项目背景与需求", "行业痛点、角色需求、设计目标与当前证据"),
            ("二、总体方案与系统架构", "端-边-云架构、技术栈、单文件与源码树交付"),
            ("三、已实现功能与模块", "12 个模块矩阵、重点模块、目录去重与边界"),
            ("四、关键技术与创新", "可信告警、SSO、本地推理、数据保护、局域网接入"),
            ("五、典型业务流程", "登录、机位配置、告警闭环、火检、跨设备上传"),
            ("六、工程实现与测试验证", "代码组织、真实自检结果、部署与限制"),
            ("七、应用场景与推广", "普通/特殊考场、保密室、安检、教学、区域"),
            ("八、安全、合规与伦理", "最小必要、同意、审计、删除、人工兜底"),
            ("九、团队分工与实施计划", "角色分工、里程碑、资源与风险"),
            ("十、结论", "项目价值、边界与下一步"),
            ("附录", "模块入口、运行方式、术语、版本变更"),
        ],
        widths=[5.2, 11.8], font_size=10,
    )
    page_break()


# ---------------------------------------------------------------------------
# 正文
# ---------------------------------------------------------------------------
def section_background():
    h1("一、项目背景与需求")
    h2("1.1 背景")
    body("教育考试考场已普遍具备视频监控、拾音、门禁与网络条件，但“有设备”不等于“有监管能力”。"
         "传统巡考高度依赖人工值守与事后翻录像：异常发现滞后、判定标准不一、取证成本高；"
         "考场设备在非考试期长期闲置，资产利用率低；全量视频上云又带来带宽、隐私与合规压力。"
         "本项目不追求替换既有硬件，而是以软件把已有设备组织成“可发现、可研判、可处置、可审计”的巡考能力。")
    h2("1.2 核心痛点")
    for t in [
        "**发现难**：人工长时间盯屏易疲劳，考场异常行为漏检；违规往往只能考后回看录像。",
        "**判定难**：不同监考人员尺度不一致，缺少可解释的判定依据与统一留痕。",
        "**处置慢**：告警从发现到终端确认、复核、归档链路长，跨设备/跨模块协同成本高。",
        "**利用低**：监控设备仅在考试期使用，日常教学与安全场景复用不足。",
        "**合规压力**：考生人脸、证件、视频属敏感个人信息，全量上云与长期保存风险高。",
    ]:
        bullet(t)

    h2("1.3 用户需求与系统响应")
    table(
        ["角色", "核心诉求", "系统响应"],
        [
            ("管理者", "全域态势、事件可追溯、操作可审计、合规达标", "总控台统一入口、驾驶舱、审计链、数据保护与验收报告"),
            ("巡考员", "减少无效巡查、快速定位异常、闭环处置", "02 机位巡检/工单、03 告警分级与复核、11 跨设备上传"),
            ("监考员", "告警轻量、不干扰考场、操作简单", "03 终端分级提示、强制确认、断网补传、角色权限"),
            ("考生", "公平、隐私、特殊需求有人工通道", "09 无障碍替代核验、05 最小必要与脱敏、全流程留痕"),
            ("运维/安保", "设备在线、画面可信、环境安全", "02 画质诊断、06 安检、07 频谱、08 保密室/火检"),
        ],
        widths=[2.8, 6.4, 8.0], font_size=9,
    )
    caption("表 1  不同角色的核心需求与系统响应")

    h2("1.4 设计目标与当前证据")
    table(
        ["目标", "设计口径", "当前证据 / 状态"],
        [
            ("异常发现", "行为、画质、安检、频谱、火情多源发现", "02/06/07/08 已实现；真实摄像头/传感器需现场联调"),
            ("响应时延", "P95 告警响应 ≤ 3 秒（本地链路）", "本地子进程/HTTP 链路已通；现场网络时延待实测"),
            ("判定质量", "AI 提示 + 人工最终判定，误报可回流", "03 强制确认/复核/误报统计已实现"),
            ("数据合规", "最小必要、脱敏、加密、可审计、可销毁", "05 已实现分级、SM3/SM2 完整性、审计与销毁台账"),
            ("交付可用", "单文件 exe + 源码树，均可离线运行", "exe 约 511 MB；源码树各模块可独立运行"),
            ("测试可复现", "每模块有自检/测试入口", "02 108/108、08 114/114、00 39/39、06 152/152 等"),
        ],
        widths=[2.8, 7.0, 7.4], font_size=9,
    )
    caption("表 2  设计目标与当前实现证据")


def section_architecture():
    h1("二、总体方案与系统架构")
    h2("2.1 设计原则")
    for t in [
        "**本地优先**：音视频在考点本地处理，原始视频不上云，只上传告警与结构化数据。",
        "**人机协同**：AI 只做提示与分级，红色告警必须人工确认，复核结论回流算法迭代。",
        "**模块解耦**：各专题独立进程/独立数据库，通过 00 联动总线与 03 终端对接，单模块故障不拖垮全局。",
        "**一次登录**：总控台统一认证，签名令牌直达模块，避免重复输入口令。",
        "**合规内嵌**：最小必要、脱敏、加密、审计、人工兜底写进默认流程，而非事后补丁。",
    ]:
        bullet(t)

    h2("2.2 端-边-云协同架构")
    body("系统按“端侧采集—边缘侧推理—业务侧处置”组织。端侧复用考场摄像机、拾音器、安检门、"
         "门禁与环境传感器；边缘侧在考点本地运行 01/02/06/07/08/10 等模块，完成人脸核验、画质诊断、"
         "行为识别、安检判据、频谱监测、保密室监控与火场识别；业务侧由总控台、03 智能终端、05 数据保护、"
         "11 局域网上传与 00 联动总线组成，负责告警汇聚、人工复核、权限与审计。")
    if os.path.isfile(ARCH):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(ARCH, width=Cm(16.4))
        caption("图 1  智能考场 AI 巡考系统端-边-云协同架构")
    body("在部署上，系统支持三种形态：① 单机演示/小型考点：一台电脑运行总控台与全部模块；"
         "② 考点局域网：边缘机运行 01/02/06/07/08/10，业务机运行 03/05/11，通过 HTTP 与共享目录对接；"
         "③ 区域平台：各考点只上报告警与结构化数据，由区域侧汇聚（接口预留）。")

    h2("2.3 技术栈与关键选型")
    table(
        ["层次", "技术/组件", "选型理由"],
        [
            ("桌面与界面", "Python 3.12 + tkinter/ttk", "零额外运行时依赖、离线可用、易于单文件打包"),
            ("视觉与AI", "OpenCV、YOLO11n ONNX、可选 VLM（Ollama/OpenAI 兼容）", "本地推理为主；模型缺失时按真实状态降级"),
            ("数据", "SQLite（WAL）+ 哈希链审计", "单机可靠、易部署；审计链可验证"),
            ("安全", "PBKDF2 口令哈希、SM3/SM2 完整性、DPAPI 密钥环", "口令不落明文；数据篡改可检出"),
            ("联动", "子进程 + JSON、HTTP（局域网）", "避免 config/models 同名冲突；跨机可部署"),
            ("交付", "自定义单文件启动器 + payload.zip + 一键打包脚本", "双击即用；图标/资源可复现构建"),
        ],
        widths=[2.6, 7.2, 7.4], font_size=9,
    )
    caption("表 3  主要技术栈与选型")

    h2("2.4 单文件交付与源码树")
    body("成品采用“启动器 + payload + 尾部索引”的单文件结构：首次运行解压到本地应用目录，"
         "之后直接复用；启动器在 `%LOCALAPPDATA%` 不可写时自动回退 `%TEMP%` 或 exe 同目录，"
         "不删除旧版本目录。源码树（项目树/All）保留每个模块的独立入口，"
         "无 runtime 时总控台自动使用系统 Python；每个模块也可单独运行，便于答辩现场排障与二次开发。")


def section_functions():
    h1("三、已实现功能与模块")
    h2("3.1 总控台：统一入口与一次登录")
    body("总控台是系统的统一入口：启动时完成一次账户/口令认证，之后进入任一模块不再重复输入。"
         "登录成功后，总控台签发带过期时间的 HMAC 令牌，模块校验令牌、模块号与签名后直接建立本模块会话；"
         "直接双击模块自己的入口时，模块仍保留原有登录页，避免绕过入口门禁。"
         "界面统一为深色主题，模块树按业务分组，右侧显示模块说明、入口、角色与运行状态。")
    for t in [
        "账户与角色：owner / admin / patrol1 / inv01 / check1，按角色限制可进入模块；",
        "审计：登录、进入模块、账户管理均写入 `总控台审计.jsonl`；",
        "模块直达：01/02/03/05/08 读取 SSO 令牌自动登录，减少现场重复输入；",
        "交付：单文件 exe 与源码树均可运行；源码树各模块可独立启动。",
    ]:
        bullet(t)

    h2("3.2 模块矩阵")
    table(
        ["序号", "模块", "已实现能力（摘要）", "验证入口 / 结果"],
        [
            ("00", "集成联动", "9 个专题的子进程 JSON 联动；通用 call/status；局域网记录上传", "39/39；LAN 38/38；HTTP 17/17"),
            ("01", "人脸识别与身份核对", "1:1/1:N、证件要素、活体、审计、网页端账户", "186 项测试 + 自检"),
            ("02", "摄像头核查与画质诊断", "机位台账/新增编辑、基准图、三级巡检、工单、报告、YOLO/VLM", "260 单元 + 108/108"),
            ("03", "智能终端操作", "告警分级/确认/复核、驾驶舱、命令台、数据维护、跨机账户", "308 单元 + 78/78 + 21/21 + 26/26"),
            ("04", "巡考机器人", "方案文档（路径规划/硬件选型/应用场景）", "方案文档，未提供可运行入口"),
            ("05", "考生数据保护", "分级、SM3/SM2 完整性、加密、脱敏、销毁、审计、验收", "174 单元 + 110 e2e + 16 GUI + 6 指标"),
            ("06", "智能安检", "安检门判据、复检、手机保管柜、验收报告", "152/152"),
            ("07", "无线信号监测", "频谱扫描、异常定位、屏蔽管控、合规与验收", "375 项自检"),
            ("08", "试卷流转/保密室", "试卷流转、保密室监控、门禁/环境告警、火场识别", "114/114；火检 17/17"),
            ("09", "无障碍与特殊考生", "通道台账、特殊档案、替代核验、公平性、演练、验收", "自检通过"),
            ("10", "普通/特殊考场", "一套平台多种模板、准入自检、分级告警、数据分域、KPI", "210 项 + 界面 36 项"),
            ("11", "告警上传与终端处置", "本机上传 03；0.0.0.0:5030 局域网 HTTP 上传；03 告警回显", "20 项 + LAN 实测"),
        ],
        widths=[1.2, 3.0, 8.3, 5.0], font_size=8.5, align_center=[0],
    )
    caption("表 4  已实现模块与验证证据（2026-10-20）")

    h2("3.3 重点模块")
    h3("（1）02 摄像头核查与画质诊断")
    body("02 是巡考发现的主入口。设备台账支持新增、编辑、删除、采集基准图；"
         "巡检按“设备在线—画面质量—视野/时钟/音频”分级给出缺陷与整改工单；"
         "实时监控支持多路画面、逐人行为识别、作弊时序识别，可选多模态大模型判读；"
         "画质可信度会随告警一起上报 03，低可信度红色告警自动降级，防止画质问题造成误报堆积。"
         "针对现场反馈的“机位配置不便”，本版把保存按钮固定在弹窗底部，补安装位置、考场必填、"
         "视场角校验与保存异常提示，并支持双击编辑，避免窗口高度不足导致“保存”按钮不可见。")
    h3("（2）03 智能终端操作")
    body("03 是告警处置的作业入口。告警按红/橙分级置顶，红色告警必须确认后才能复核；"
         "复核要求结论、理由、见证人与附件，全部写入哈希链审计。驾驶舱聚合设备、告警、处置与考场态势；"
         "命令台可向 01/02 下发命令并查看回执；数据维护支持按权限清空并留痕。"
         "本版重做告警处置页：深色表格、级别文字上色/加粗、统计条、选中详情与双击确认，"
         "解决原稿截图中“整行红底、文字对比度低”的问题。")
    h3("（3）05 考生数据保护")
    body("05 把“最小必要、目的限定、加密、脱敏、可审计、可销毁”落到代码："
         "账户口令 PBKDF2 加盐；敏感字段加密；SM3 哈希 + SM2 签名保证已加密数据不可被静默修改；"
         "查看明文、共享、销毁等高风险管理动作需二次确认并留痕；导出验收报告与六项 KPI。"
         "密钥环使用 Windows DPAPI，打包时不携带旧密钥，首次运行按当前 Windows 用户重建，"
         "避免“换机后密钥不可用”的弹窗。")
    h3("（4）06/07 安检与无线信号")
    body("06 以安检门判据引擎为核心，串起过检录入、报警复检、手机保管柜、特殊考生通道与验收报告；"
         "07 覆盖频谱扫描、异常定位、屏蔽设备开关、考试时间表与合规验收，支持采集服务与一键演示。"
         "两者均把“设备报警”转成结构化事件，可经 00 联动进入 03。")
    h3("（5）08 试卷流转、保密室与火场识别")
    body("08 负责试卷袋流转、保密室多路监控、门禁/红外/水浸/烟感等告警与 KPI。"
         "本版把原 13 的火场识别整体并入 08：实时画面页新增“火场识别”开关，"
         "多路画面后台分析火焰/烟雾，连续帧确认后以 FIRE_ALARM（红）/FIRE_SMOKE（橙）进入告警中心，"
         "不阻塞取流；火检功能冒烟 17/17。")
    h3("（6）09/10 特殊考生与考场模板")
    body("09 为特殊考生保留人工替代核验、受阻事件登记、公平性分层测试与演练，强调“人工兜底不可取消”；"
         "10 把普通考场与特殊七类考场的配置差异、准入判据、分级告警、数据分域与验收指标做成模板，"
         "支持一套平台多套模板。")
    h3("（7）11 告警上传与局域网接入")
    body("11 提供桌面 UI：本机填表即可把告警写入 03；开启局域网接收后，"
         "其他电脑或手机浏览器访问 http://<本机IP>:5030/ 即可上传，上传结果与 03 最近告警同页回显。"
         "服务支持可选访问令牌，适合答辩现场多设备演示，也适合考点内多台终端向业务机汇聚告警。")
    h3("（8）00 集成联动")
    body("00 以“子进程 + JSON”隔离各专题的同名 config/models，提供统一的 call(key, op, args) 与 "
         "status()；当前接入 01/02/03/05/06/07/08/09/10 共 9 个专题，主链覆盖"
         "“02 画质 → 03 告警降权 → 01 核验 → 03 复核 → 02 工单 → 06 安检 → 06 复检”。")

    h2("3.4 目录去重与边界说明")
    body("原目录中的 11 摄像头 AI 识别与 13 火场识别属于重复能力：13 的火检已并入 08；"
         "11 的逐人框选/行为识别已由 02 的 person_detect 覆盖。本版删除这两个目录，"
         "总控台结构目录重排为 00–11（04 为方案文档），避免评委与用户看到重复入口。")
    body("需要如实说明的边界：04 巡考机器人目前只有方案文档，没有硬件与可运行入口；"
         "07 的频谱/屏蔽、08 的门禁/环境传感、10 的机位接入均提供真实接口与模拟数据，"
         "但现场精度、时延与合规效果仍需在真实考点联调后复测；01/02 的识别指标依赖模型与现场素材，"
         "项目书不把实验室结果等同于现场验收。")


def section_tech():
    h1("四、关键技术与创新")
    table(
        ["创新点", "技术做法", "解决的问题"],
        [
            ("分级告警 + 人工最终判定", "红/橙两级、强制确认、复核结论/理由/见证人/附件、哈希链留痕", "AI 误报与责任边界；处置过程可追溯"),
            ("画质可信度降权", "02 把机位缺陷折算为 credibility；03 收到低可信度红色告警自动降橙", "画质差导致的误报堆积"),
            ("本地边缘推理 + 可选 VLM", "OpenCV/YOLO11n 本地推理；可选 Ollama/OpenAI 兼容 VLM 判读", "带宽/隐私/断网；复杂场景补充判读"),
            ("统一登录 + HMAC SSO", "总控台一次认证，签发短时令牌，模块校验签名/模块号/过期时间", "重复输入口令；模块各自认证割裂"),
            ("模块化独立运行", "各专题独立进程/数据库；包根路径引导；无 runtime 回退系统 Python", "单模块故障拖垮全局；现场排障困难"),
            ("跨设备局域网上传", "11 内置 HTTP 服务与网页表单，手机/电脑直接上传 03", "不同设备告警汇聚；答辩演示多端协同"),
            ("数据保护内嵌", "PBKDF2、SM3/SM2、DPAPI、脱敏、销毁台账、审计", "考生敏感信息全生命周期合规"),
            ("单文件交付", "启动器 + payload + 尾部索引；图标/资源嵌入；多级解压回退", "现场部署与双击即用"),
        ],
        widths=[3.6, 8.2, 5.4], font_size=8.8,
    )
    caption("表 5  关键技术与创新点")

    body("其中，最能与同类方案形成差异的是“可信告警”与“可运行交付”两点："
         "前者用画质可信度、分级确认、复核留痕把 AI 从“黑盒判断”变成“可解释的线索”；"
         "后者把系统做成单文件 exe 与可独立运行的源码树，答辩现场不依赖云、不依赖外网，"
         "每个模块都能单独启动、单独自检。")


def section_workflow():
    h1("五、典型业务流程")
    h2("5.1 统一登录与模块进入")
    for i, t in enumerate([
        "启动总控台，输入一次账户/口令，选择“统一登录”；",
        "总控台校验账户、角色与可进入模块，写入审计；",
        "双击模块，总控台签发 SSO 令牌并启动模块；",
        "模块校验令牌后直接建立会话；用户不再输入第二遍口令；",
        "需要换人时点“切换账号”，重新登录一次。",
    ], 1):
        body("%d. %s" % (i, t), indent=False, space_after=2)

    h2("5.2 摄像头机位配置与巡检")
    for i, t in enumerate([
        "在 02「设备台账」点「新增机位」，填写设备编码、考场、取流地址（0/1/RTSP/本地文件）等；",
        "保存后可在台账看到机位；双击或「编辑机位」修改，基准图/心跳状态保留；",
        "采集基准图，执行三级巡检，查看逐机位缺陷与整改工单；",
        "工单按“发现—派单—整改—复核”闭环，整改后截图与复核意见留痕；",
        "告警上报 03 时携带画质可信度，低可信度红色告警自动降级。",
    ], 1):
        body("%d. %s" % (i, t), indent=False, space_after=2)

    h2("5.3 告警处置闭环")
    for i, t in enumerate([
        "02/06/08/11 产生告警并上传 03；",
        "03 按红/橙分级展示，红色置顶并提示强制确认；",
        "值班人员确认告警，填写处置说明；",
        "复核人员选择结论（成立/不成立/误报），填写理由、见证人与附件；",
        "系统写入哈希链审计，统计处置时延与误报率，回流算法迭代。",
    ], 1):
        body("%d. %s" % (i, t), indent=False, space_after=2)

    h2("5.4 火场识别流程")
    for i, t in enumerate([
        "08 实时画面页开启「火场识别」；",
        "多路画面后台按固定间隔分析火焰颜色/闪烁/增长/烟雾等信号；",
        "连续帧确认后生成火警事件，写入告警中心（红/橙）；",
        "值班人员在 03 或 08 告警中心确认、处置并留痕；",
        "原始视频本地留存，仅告警与结构化数据进入业务层。",
    ], 1):
        body("%d. %s" % (i, t), indent=False, space_after=2)

    h2("5.5 跨设备局域网上传")
    for i, t in enumerate([
        "在 11「局域网接收」页设置监听地址/端口（默认 0.0.0.0:5030），点「启动接收」；",
        "其他电脑/手机连接同一局域网，浏览器打开 http://<本机IP>:5030/；",
        "填写告警码、设备、考场、详情，提交；",
        "11 调用 03 的接入接口写入终端，页面回显结果与 03 最近告警；",
        "03 端继续确认/复核，形成闭环。",
    ], 1):
        body("%d. %s" % (i, t), indent=False, space_after=2)


def section_engineering():
    h1("六、工程实现与测试验证")
    h2("6.1 代码组织")
    body("项目树按专题分目录：00 集成联动、01 人脸、02 摄像头、03 终端、04 机器人方案、"
         "05 数据保护、06 安检、07 无线信号、08 试卷/保密室/火检、09 无障碍、10 考场模板、"
         "11 告警上传。每个目录包含入口脚本、自检/测试、说明文档与独立数据目录；"
         "公共的总控台、账户、SSO、主题与打包工具位于项目根目录。")
    h2("6.2 测试与自检")
    body("项目坚持“可复现证据”：每个模块提供自检/单元测试/端到端脚本，"
         "关键结果如下（2026-10-20 实测）：")
    table(
        ["模块", "测试/自检结果", "说明"],
        [
            ("00 集成联动", "39/39；LAN 38/38；发布命令 HTTP 17/17", "9 个专题 status 全部可调用"),
            ("01 人脸", "186 项测试 + 自检", "1:1/1:N、证件要素、审计"),
            ("02 摄像头", "260 单元 + 108/108 自检", "含机位配置、巡检、工单、报告"),
            ("03 终端", "308 单元 + 78/78 + 21/21 + 26/26", "含驾驶舱、命令台、LAN 服务"),
            ("05 数据保护", "174 单元 + 110 e2e + 16 GUI + 6 指标", "含 SM3/SM2、脱敏、销毁"),
            ("06 安检", "152/152", "零第三方依赖"),
            ("07 无线信号", "375 项自检", "频谱/屏蔽/合规/联动"),
            ("08 试卷/保密室", "114/114；火检 17/17", "含门禁/环境/火场识别"),
            ("09 无障碍", "自检通过", "通道/档案/替代核验/演练"),
            ("10 考场模板", "210 项 + 界面 36 项", "普通 + 特殊七类模板"),
            ("11 告警上传", "20 项 + LAN 实测", "HTTP 上传/回显"),
        ],
        widths=[3.2, 8.4, 5.6], font_size=8.8,
    )
    caption("表 6  模块测试与自检结果")

    h2("6.3 运行与部署")
    for t in [
        "**单文件 exe**：双击运行，首次解压到本地应用目录；图标与资源已嵌入，约 511 MB；",
        "**源码树**：python 总控台.py 启动；各模块目录可直接运行；需要依赖时按各目录说明安装；",
        "**离线**：模型/依赖随包或本地安装；断网时本地推理、缓存与补传仍可用；",
        "**局域网**：11 提供告警上传入口；03 提供局域网接入服务；01/02 可上报操作记录。",
    ]:
        bullet(t)

    h2("6.4 真实性与限制")
    body("项目书中的“已实现”以代码入口与自检结果为准；不把方案文档写成已实现，"
         "不把实验室合成样本的指标写成现场验收结论。巡考机器人、真实频谱仪/门禁/摄像头、"
         "区域平台对接仍需现场联调；05 的密钥环与 Windows 用户绑定，换机需按说明重建密钥；"
         "AI 识别受光照、遮挡、模型版本影响，必须保留人工复核与替代通道。")


def section_application():
    h1("七、应用场景与推广")
    h2("7.1 典型场景")
    table(
        ["场景", "部署方式", "主要价值"],
        [
            ("标准化考场", "单机/考点局域网，利旧摄像机与拾音器", "行为识别、画质核查、分级告警、闭环处置"),
            ("特殊考场（七类）", "10 模板 + 09 无障碍替代核验", "差异化配置、人工兜底、公平性指标"),
            ("试卷保密室", "08 保密室监控 + 门禁/环境/火检", "六面覆盖、开柜/开门告警、火情早期发现"),
            ("智能安检入口", "06 安检门/手持探测器/保管柜", "手机/无线耳机/金属器材事前阻断"),
            ("无线信号管控", "07 频谱监测 + 屏蔽设备", "异常信号发现、定位、按时间表管控"),
            ("日常教学/巡课", "双模式切换，数据隔离", "设备复用、课堂行为分析、教学督导"),
            ("区域/多考点", "各考点只上报告警与结构化数据", "区域态势、统一审计、分级运维"),
        ],
        widths=[3.4, 7.0, 6.8], font_size=8.8,
    )
    caption("表 7  典型应用场景")

    h2("7.2 推广路径")
    body("第一阶段在单考点完成“02 + 03 + 05”最小闭环，验证画质、告警、处置与合规；"
         "第二阶段接入 06/07/08/09/10，形成入口—考场—保密室—特殊考生的全场景；"
         "第三阶段在区域平台汇聚多考点数据，输出验收报告与运维指标。"
         "系统以利旧为主，硬件增量集中在安检、频谱、门禁等必要环节，具体投入以试点测算为准。")

    h2("7.3 社会与育人价值")
    body("项目把“AI 不替代人、而是让人更快发现和更规范处置”作为设计底线："
         "所有智能环节都保留人工复核与替代通道；特殊考生不会被设备门槛挡在门外；"
         "考生数据按最小必要处理并可审计、可销毁。对学校而言，系统同时服务考试安全与日常教学，"
         "提高设备利用率；对学生团队而言，项目覆盖视觉、后端、桌面、数据安全、测试与文档的完整工程链。")


def section_security():
    h1("八、安全、合规与伦理")
    for t in [
        "**最小必要**：只采集考试安全必需的音视频/身份要素；原始视频本地留存，仅告警与结构化数据上传。",
        "**单独同意**：人脸信息处理取得考生单独同意，不捆绑一般考试规则；保留非人脸替代通道。",
        "**目的限定**：数据仅用于考试安全与作弊判定，不做情绪分析、行为画像等超范围使用。",
        "**可审计**：登录、告警、复核、查看明文、导出、销毁等关键动作写入审计链，可校验完整性。",
        "**可删除**：按约定保留期到期删除/销毁，销毁记录留痕；05 提供销毁台账与验收报告。",
        "**人工兜底**：09 保留人工核验、双人见证、受阻事件登记，避免设备故障造成新的不公平。",
        "**风险提示**：AI 识别存在误报/漏报，系统以“提示 + 人工判定”呈现，不自动作出处分结论。",
    ]:
        bullet(t)
    body("技术措施与管理措施需配套：建议学校明确数据责任人、保留期限、访问审批与应急流程；"
         "跨校/区域部署时，应在合同与告知中写清数据边界，并按当地考试管理机构要求进行安全评估。")


def section_team():
    h1("九、团队分工与实施计划")
    h2("9.1 团队分工（待填）")
    table(
        ["角色", "主要职责", "成员（待填）"],
        [
            ("项目负责人", "总体方案、进度管理、答辩与文档统稿", "【待填】"),
            ("视觉/算法", "人脸、行为识别、画质诊断、火检参数与测试", "【待填】"),
            ("后端/数据", "终端、告警、审计、数据保护、联动总线", "【待填】"),
            ("桌面/交互", "总控台、模块 UI、深色主题、可用性测试", "【待填】"),
            ("测试/交付", "自检脚本、打包、部署、现场联调", "【待填】"),
        ],
        widths=[3.2, 9.4, 4.6], font_size=9,
    )
    caption("表 8  团队分工（请按实际填写）")

    h2("9.2 实施计划")
    table(
        ["阶段", "时间", "目标", "状态"],
        [
            ("原型与模块实现", "2026 年上半年", "完成 01/02/03/05/06 等核心模块与自检", "已完成"),
            ("整体集成与 UI", "2026 年 10 月", "总控台统一登录、SSO、深色主题、模块去重", "已完成"),
            ("现场联调与试点", "2026 年 11 月起", "真实摄像头/门禁/频谱/安检设备联调", "待开展"),
            ("区域推广", "后续", "多考点数据汇聚、验收与运维体系", "规划中"),
        ],
        widths=[3.6, 4.4, 7.2, 2.6], font_size=9, align_center=[3],
    )
    caption("表 9  实施里程碑")

    h2("9.3 资源与风险")
    body("开发环境为普通 Windows 笔记本即可；演示不需要专有云资源。主要风险是现场设备协议差异、"
         "模型对光照/遮挡的敏感性、以及跨机部署的网络与权限配置。应对方式是保留模拟数据与"
         "真实接口双通道、提供自检与日志、先小范围试点再推广。")


def section_conclusion():
    h1("十、结论")
    body("本项目把“端-边-云协同、分级告警、人工复核、数据保护”从方案落成可运行的软件系统："
         "12 个模块覆盖人脸、画质、终端、数据保护、安检、频谱、保密室/火检、无障碍、考场模板与"
         "跨设备上传；总控台统一登录一次即可直达模块；单文件 exe 与源码树两种交付形态均可在离线环境运行。")
    body("项目的核心价值不是“用 AI 替代监考”，而是把 AI 变成可解释、可复核、可追溯的线索，"
         "让巡考人员更快发现、更规范处置，让考生数据得到最小必要与全流程保护。"
         "下一步将重点完成真实设备联调、现场指标复测与区域平台对接，并在试点中持续迭代算法与流程。")


def appendix():
    h1("附录")
    h2("附录 A  模块与入口")
    table(
        ["序号", "目录", "入口", "说明"],
        [
            ("00", "00_集成联动", "main.py", "联动总线/命令执行器/局域网记录上传"),
            ("01", "01_人脸识别与身份核对/01_人脸识别与身份核对", "gui.py", "人脸核验 GUI + 网页端"),
            ("02", "02_摄像头核查与视频质量诊断", "gui.py", "机位台账/巡检/工单/实时监控"),
            ("03", "03_智能终端操作", "terminal_gui.py", "告警处置/复核/驾驶舱/命令台"),
            ("04", "04_巡考机器人", "—", "方案文档（未提供可运行入口）"),
            ("05", "05_考生数据保护", "gui.py", "数据分级/加密/脱敏/审计/验收"),
            ("06", "06_智能安检与手机作弊器材检测", "app.py", "安检判据/复检/保管柜"),
            ("07", "07_无线信号监测与屏蔽管控", "gui.py", "频谱/定位/屏蔽/合规"),
            ("08", "08_试卷流转与保密室智能监控", "gui.py", "试卷流转/保密室/火检"),
            ("09", "09_无障碍与特殊考生保障", "app.py", "通道/档案/替代核验/演练"),
            ("10", "10_普通考场+特殊考场应用", "launcher.py", "考场模板/准入/告警/分域"),
            ("11", "11_告警上传与终端处置演示", "gui.py", "本机上传/局域网 HTTP/回显"),
        ],
        widths=[1.2, 6.4, 2.8, 6.6], font_size=8.6, align_center=[0],
    )
    caption("表 10  模块目录与入口")

    h2("附录 B  运行方式")
    body("源码树：进入 项目树/All，执行 python 总控台.py；各模块目录也可直接运行。"
         "单文件：双击 智能考场AI巡考系统.exe，首次解压后自动打开总控台。"
         "一键打包：在 exe打包文件 目录用 Git Bash 执行 bash 一键打包.sh。")
    body("默认账户见项目 README；首次登录后请修改口令。项目书不列出具体口令，避免公开传播。")

    h2("附录 C  术语")
    table(
        ["术语", "解释"],
        [
            ("端-边-云", "端侧采集、边缘侧本地推理、业务侧汇聚与处置的三层架构"),
            ("SSO", "单点登录；本项目用 HMAC 签名令牌让模块免二次输入口令"),
            ("分级告警", "红/橙两级；红色需强制确认，橙色待复核"),
            ("画质可信度", "02 根据机位缺陷折算的 0~1 系数，用于 03 告警降权"),
            ("哈希链审计", "每条记录包含前一条哈希，任何篡改可被定位"),
            ("DPAPI", "Windows 数据保护接口，用于本机密钥环加密"),
        ],
        widths=[3.2, 13.8], font_size=9,
    )
    caption("表 11  术语简释")

    h2("附录 D  版本变更")
    body("V1.0：原始方案框架（端-边-云、场景、指标设想）。")
    body("V2.0（2026-10-20）：按实际项目重写；补充已实现模块、统一登录/SSO、深色 UI、"
         "02 机位配置修复、03 告警 UI、11 局域网上传、08 火检、05 密钥环处理；"
         "删除重复目录 11/13 并重排结构目录；补充真实自检结果与边界说明。")


def main():
    make_arch_image(ARCH)
    init_doc()
    cover()
    declaration()
    abstract()
    revision_note()
    toc()
    section_background()
    section_architecture()
    section_functions()
    section_tech()
    section_workflow()
    section_engineering()
    section_application()
    section_security()
    section_team()
    section_conclusion()
    appendix()
    doc.save(OUT)
    print("saved", OUT, os.path.getsize(OUT))


if __name__ == "__main__":
    main()
