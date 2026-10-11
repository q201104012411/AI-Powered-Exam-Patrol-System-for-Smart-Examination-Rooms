import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "candidates"

# 表头
# 说明：
#   id_card  —— 仅用于登记，服务端会立即加盐哈希，库里不存明文；
#   consent_status —— 单独同意登记（未登记 / 已同意 / 已拒绝 / 替代通道），
#                     未登记同意的考生在核验端会被直接拒绝（合规硬约束）。
headers = ["candidate_id", "name", "id_card", "ticket_no", "seat_no",
           "reference_photo", "consent_status"]
ws.append(headers)

# 示例数据
sample = [
    ["C001", "张三", "330100199001011234", "2026T001", "A-01", "data/ref_photos/C001.jpg", "已同意"],
    ["C002", "李四", "330100199202022345", "2026T002", "A-02", "data/ref_photos/C002.jpg", "已同意"],
    ["C003", "王五", "330100199303033456", "2026T003", "A-03", "data/ref_photos/C003.jpg", "未登记"],
]
for row in sample:
    ws.append(row)

# 填写说明页
ws2 = wb.create_sheet("填写说明")
for line in [
    "字段说明：",
    "  candidate_id     考生编号（唯一，必填）",
    "  name             姓名",
    "  id_card          身份证号（仅用于登记，服务端加盐哈希后不再保存明文）",
    "  ticket_no        准考证号（身份要素之一；常规模式下即用它核验身份，必填）",
    "  seat_no          座位号（身份要素之一，始终参与校验）",
    "  reference_photo  参考照路径（用于提取特征；建档后可删除原图）",
    "  consent_status   单独同意状态：未登记 / 已同意 / 已拒绝 / 替代通道",
    "",
    "注意：",
    "  1. 未登记“已同意”或“替代通道”的考生，核验端会直接拒绝（合规要求）；",
    "  2. 参考照请使用清晰单人正脸照；建档（enroll.py build）后可执行 --purge-photos 删除原图；",
    "  3. 演示数据可用编造号码；生产环境请开启 STRICT_ID_CHECKSUM=1 校验身份证校验位。",
]:
    ws2.append([line])
ws2.column_dimensions["A"].width = 90

# 样式
header_fill = PatternFill("solid", fgColor="2D6CDF")
header_font = Font(color="FFFFFF", bold=True, size=11)
center = Alignment(horizontal="center", vertical="center")
thin = Side(style="thin", color="D0D7E5")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col_idx, _ in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center
    cell.border = border

# 列宽
widths = [14, 12, 22, 14, 10, 32, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

# 数据行边框
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical="center")

# 冻结首行
ws.freeze_panes = "A2"

wb.save("candidates_template.xlsx")
print("已生成：candidates_template.xlsx")